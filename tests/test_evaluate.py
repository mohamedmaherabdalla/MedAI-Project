import subprocess
import sys

import numpy as np
import pytest
from openpyxl import load_workbook
from PIL import Image

from medai.evaluate import RUN_COLUMNS, add_log_line, add_note, add_run, evaluate, log_facts, results_lock, seed_note


def save(folder, name, pixels):
    folder.mkdir(parents=True, exist_ok=True)
    Image.fromarray(np.array(pixels, dtype=np.uint8)).save(folder / f"{name}.png")


def make_example(gt_dir, pred_dir):
    # Frame a: 4 true pixels, 4 predicted, 2 shared -> Dice 4/8 = 0.5, IoU 2/6.
    # 127 is below 0.5 after /255 and must count as background; 200 counts as polyp.
    save(gt_dir, "a", [[255, 255, 0], [255, 255, 0], [0, 0, 0]])
    save(pred_dir, "a", [[200, 255, 255], [127, 0, 255], [0, 0, 0]])
    # Frame b: prediction equals the mask -> Dice 1, IoU 1.
    save(gt_dir, "b", [[0, 255], [0, 0]])
    save(pred_dir, "b", [[0, 255], [0, 0]])
    # Frame c: polyp missed entirely -> Dice 0, IoU 0, one empty prediction.
    save(gt_dir, "c", [[255, 0], [0, 0]])
    save(pred_dir, "c", [[127, 0], [0, 0]])


def test_mean_dice_and_iou_on_a_known_example(tmp_path):
    make_example(tmp_path / "gt", tmp_path / "pred")
    dice, iou, empty = evaluate(tmp_path / "pred", tmp_path / "gt")
    assert dice == pytest.approx((0.5 + 1 + 0) / 3)
    assert iou == pytest.approx((2 / 6 + 1 + 0) / 3)
    assert empty == 1


def test_refuses_c6_without_final(tmp_path):
    gt_dir, pred_dir = tmp_path / "data_C6" / "masks_C6", tmp_path / "pred"
    make_example(gt_dir, pred_dir)
    with pytest.raises(SystemExit):
        evaluate(pred_dir, gt_dir)
    assert evaluate(pred_dir, gt_dir, final=True)[0] == pytest.approx(0.5)


def test_results_sheet_has_runs_and_notes(tmp_path):
    path = tmp_path / "results.xlsx"
    add_run({column: column for column in RUN_COLUMNS}, path)
    add_note("2026-10-14", "someone", "a note", path)
    book = load_workbook(path)
    assert [c.value for c in book["runs"][2]] == RUN_COLUMNS
    assert [c.value for c in book["notes"][2]] == ["2026-10-14", "someone", "a note"]
    with pytest.raises(KeyError):
        add_run({"run_id": "incomplete"}, path)


def test_log_facts_reads_the_job_log(tmp_path):
    log = tmp_path / "slurm-x-17.log"
    log.write_text("job 17 on ws-l1-006, config /c/x.yaml, code 145673b8863d2f656f30719de9f597647541982d\n"
                   "name, memory.used [MiB], memory.total [MiB]\nNVIDIA RTX 5000 Ada Generation, 182 MiB, 32760 MiB\n"
                   "RunTime=1-02:31:27\nStartTime=2026-10-14T16:48:05\n")
    assert log_facts([log]) == ("NVIDIA RTX 5000 Ada Generation", "2026-10-14", 26.52, "145673b", "17")


def test_seed_note_once_seeds_0_1_2_are_in(tmp_path):
    path = tmp_path / "results.xlsx"
    for seed, dice in [(0, 0.5), (1, 0.6)]:
        add_run(dict({c: 1 for c in RUN_COLUMNS}, run_id=f"m_c1-c4_seed{seed}", model="M", seed=seed, Dice=dice), path)
        assert seed_note(f"m_c1-c4_seed{seed}", path) is None
    add_run(dict({c: 1 for c in RUN_COLUMNS}, run_id="other_c1-c4_seed2", seed=2), path)
    assert seed_note("other_c1-c4_seed2", path) is None
    add_run(dict({c: 1 for c in RUN_COLUMNS}, run_id="m_c1-c4_seed2", model="M", seed=2, Dice=0.7), path)
    note = seed_note("m_c1-c4_seed2", path)
    assert "Dice 0.6000 +- 0.1000" in note and "IoU 1.0000 +- 0.0000" in note


def test_log_line(tmp_path):
    row = {"date": "2026-10-15", "member": "A", "model": "M", "seed": 1, "epochs": 20, "Dice": 0.5,
           "IoU": 0.25, "train time (h)": 0.3, "output path": "/out"}
    add_log_line(row, "/home/x/configs/m_seed1.yaml", "abc1234", "17", tmp_path / "log.md")
    assert (tmp_path / "log.md").read_text() == ("| 2026-10-15 | A | configs/m_seed1.yaml | abc1234 | M, seed 1, 20 epochs: "
                                                 "C5 Dice 0.5000, IoU 0.2500 (Slurm 17, 0.3 h) | /out |\n")


def test_results_lock_blocks_a_second_writer(tmp_path):
    probe = f"import fcntl; fcntl.lockf(open({str(tmp_path / 'results.lock')!r}, 'w'), fcntl.LOCK_EX | fcntl.LOCK_NB)"
    with results_lock(tmp_path / "results.xlsx"):
        assert subprocess.run([sys.executable, "-c", probe], capture_output=True).returncode != 0
    assert subprocess.run([sys.executable, "-c", probe], capture_output=True).returncode == 0
