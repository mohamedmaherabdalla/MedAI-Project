"""Our one evaluation: Dice and IoU per frame (masks binarised at 0.5), mean over frames.

Score one run, and with --save add its row to the "runs" sheet of docs/results.xlsx and its line to
docs/EXPERIMENT_LOG.md (each Slurm job does this itself, see scripts/train_baseline.sbatch):
    PYTHONPATH=src python -m medai.evaluate configs/<run>.yaml --save --logs <slurm log> ...
GPU type, date, train time, code commit and job id come from the Slurm logs; --date and --hours give
them for runs whose logs predate the timestamps (the sacct database was unreachable on 2026-10-14).
When seeds 0, 1 and 2 of a run are all in the sheet, a note with their mean and standard deviation is added.
"""
import argparse
import re
from contextlib import contextmanager
from datetime import date
from pathlib import Path

import numpy as np
import yaml
from openpyxl import Workbook, load_workbook
from PIL import Image

try:
    import fcntl
except ImportError:  # Windows (the PC) has no fcntl; msvcrt gives the same exclusive lock there
    fcntl = None
    import msvcrt

RESULTS = Path(__file__).resolve().parents[2] / "docs" / "results.xlsx"
EXPERIMENT_LOG = RESULTS.parent / "EXPERIMENT_LOG.md"
RUN_COLUMNS = ["run_id", "date", "member", "model", "paper (venue, year)", "repo@commit", "backbone",
               "pretrained weights", "input size", "augmentation", "loss", "optimizer", "learning rate",
               "batch size", "epochs", "checkpoint rule", "train centres", "n_train frames", "val centre",
               "n_val frames", "seed", "GPU type", "train time (h)", "Dice", "IoU",
               "empty-prediction frames on val", "notes", "output path"]
NOTE_COLUMNS = ["date", "author", "note"]
# Run columns copied straight from the config, and the config key each one comes from.
CONFIG_KEYS = {"run_id": "name", "member": "member", "model": "model", "paper (venue, year)": "paper",
               "backbone": "backbone", "pretrained weights": "pretrained_weights", "input size": "input_size",
               "augmentation": "augmentation", "loss": "loss", "optimizer": "optimizer",
               "learning rate": "learning_rate", "batch size": "batch_size", "epochs": "epochs",
               "checkpoint rule": "checkpoint_rule", "train centres": "train_centres",
               "val centre": "val_centre", "seed": "seed", "notes": "notes", "output path": "run_dir"}


def binary_mask(path):
    return np.asarray(Image.open(path).convert("L"), dtype=np.float32) / 255 > 0.5


def dice_iou(pred, gt):
    inter = np.logical_and(pred, gt).sum()
    union = np.logical_or(pred, gt).sum()
    if union == 0:
        return 1.0, 1.0  # both masks empty: the prediction is exactly right
    return 2 * inter / (pred.sum() + gt.sum()), inter / union


def evaluate(pred_dir, gt_dir, final=False):
    """Mean Dice, mean IoU and the number of empty predictions over every mask in gt_dir.

    The prediction for <frame> is pred_dir/<frame>.png.
    """
    gt_paths = sorted(Path(gt_dir).iterdir())
    if not final and any("data_C6" in str(p.resolve()) for p in gt_paths):
        raise SystemExit("C6 is the held-out test centre: evaluate on it only with --final")
    scores, empty = [], 0
    for gt_path in gt_paths:
        pred, gt = binary_mask(Path(pred_dir) / f"{gt_path.stem}.png"), binary_mask(gt_path)
        if pred.shape != gt.shape:
            raise ValueError(f"{gt_path.stem}: prediction {pred.shape} but mask {gt.shape}")
        scores.append(dice_iou(pred, gt))
        empty += int(not pred.any())
    dice, iou = np.mean(scores, axis=0)
    return float(dice), float(iou), empty


def open_results(path):
    """Load the results workbook, or create it with its 'runs' and 'notes' sheets."""
    if Path(path).exists():
        return load_workbook(path)
    book = Workbook()
    book.active.title = "runs"
    book.active.append(RUN_COLUMNS)
    book.create_sheet("notes").append(NOTE_COLUMNS)
    return book


def add_run(row, path=RESULTS):
    """Append one run to the 'runs' sheet; row is a dict with a value for every column."""
    book = open_results(path)
    book["runs"].append([row[column] for column in RUN_COLUMNS])
    book.save(path)


def add_note(date, author, note, path=RESULTS):
    book = open_results(path)
    book["notes"].append([date, author, note])
    book.save(path)


@contextmanager
def results_lock(path=RESULTS):
    """Hold an exclusive lock while writing the results, so two jobs never write the sheet at once.

    The lock file sits next to the sheet on the NFS /home, where fcntl locks hold across nodes
    (the /l Lustre mount only has node-local locks). The lock is released if the job dies.
    """
    with open(Path(path).with_suffix(".lock"), "w") as lock:
        if fcntl:
            fcntl.lockf(lock, fcntl.LOCK_EX)
        else:
            while True:  # LK_LOCK gives up after 10 s, so try again until the other writer is done
                try:
                    msvcrt.locking(lock.fileno(), msvcrt.LK_LOCK, 1)
                    break
                except OSError:
                    pass
        yield


def add_log_line(row, config, commit, jobs, path=EXPERIMENT_LOG):
    """Append the run's line to docs/EXPERIMENT_LOG.md."""
    result = (f"{row['model']}, seed {row['seed']}, {row['epochs']} epochs: C5 Dice {row['Dice']:.4f}, "
              f"IoU {row['IoU']:.4f} (Slurm {jobs}, {row['train time (h)']} h)")
    with open(path, "a") as log:
        log.write(f"| {row['date']} | {row['member']} | configs/{Path(config).name} | {commit} | {result} "
                  f"| {row['output path']} |\n")


def seed_note(run_id, path=RESULTS):
    """Mean and sample standard deviation of Dice and IoU once seeds 0, 1 and 2 of this run are all in."""
    group = run_id.rsplit("_seed", 1)[0]
    rows = [dict(zip(RUN_COLUMNS, r)) for r in open_results(path)["runs"].iter_rows(min_row=2, values_only=True)]
    rows = [r for r in rows if r["run_id"].rsplit("_seed", 1)[0] == group]
    if sorted(r["seed"] for r in rows) != [0, 1, 2]:
        return None
    dice, iou = np.array([r["Dice"] for r in rows]), np.array([r["IoU"] for r in rows])
    return (f"{rows[0]['model']} ({group}), seeds 0, 1, 2 on C5: Dice {dice.mean():.4f} +- {dice.std(ddof=1):.4f}, "
            f"IoU {iou.mean():.4f} +- {iou.std(ddof=1):.4f} (mean +- sample standard deviation over the 3 runs)")


def log_facts(logs):
    """GPU type, start date, summed Slurm run time (h), code commit and job ids from the job logs.

    Date and run time are None when a log lacks them; commit and job ids are "unknown".
    """
    gpu, start_date, hours, commit, jobs = None, None, 0.0, "unknown", []
    for text in (Path(log).read_text() for log in logs):
        found = re.search(r"^(NVIDIA [^,]+),", text, re.M)  # the nvidia-smi line
        gpu = gpu or (found and found[1])
        header = re.search(r"^job (\d+) on .*, code ([0-9a-f]{7})", text, re.M)  # first line of the log
        if header:
            jobs.append(header[1])
            commit = commit if commit != "unknown" else header[2]
        start, runtime = re.search(r"StartTime=(\S+)", text), re.search(r"RunTime=(\S+)", text)
        if not (start and runtime):
            hours = None
            continue
        start_date = start_date or start[1][:10]
        days, _, clock = runtime[1].rpartition("-")
        h, m, s = map(int, clock.split(":"))
        hours = hours if hours is None else hours + int(days or 0) * 24 + h + m / 60 + s / 3600
    return gpu, start_date, hours and round(hours, 2), commit, " + ".join(jobs) or "unknown"


def save_run(cfg, config, dice, iou, empty, logs, run_date=None, hours=None):
    """Add the run's row to the sheet and its line to the experiment log, and the seed note when due."""
    gpu, start_date, log_hours, commit, jobs = log_facts(logs)
    row = {column: cfg[key] for column, key in CONFIG_KEYS.items()}
    if "bigaug" in cfg:
        row["model"] += " + BigAug"
    row.update({"date": run_date or start_date or "unknown", "repo@commit": f"{cfg['repo_url']}@{cfg['commit']}",
                "n_train frames": len(list(Path(cfg["train_dir"], "images").iterdir())),
                "n_val frames": len(list(Path(cfg["val_dir"], "images").iterdir())),
                "GPU type": gpu or "unknown", "train time (h)": hours or log_hours or "unknown",
                "Dice": round(dice, 4), "IoU": round(iou, 4), "empty-prediction frames on val": empty})
    with results_lock():
        add_run(row)
        add_log_line(row, config, commit, jobs)
        note = seed_note(row["run_id"])
        if note:
            add_note(date.today().isoformat(), row["member"], note)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("config")
    parser.add_argument("--final", action="store_true", help="allow the held-out C6 centre")
    parser.add_argument("--save", action="store_true", help="add the run to docs/results.xlsx")
    parser.add_argument("--logs", nargs="+", default=[], help="Slurm logs of the run's jobs")
    parser.add_argument("--date", help="run date, when the logs do not have it")
    parser.add_argument("--hours", type=float, help="Slurm run time in hours, when the logs do not have it")
    parser.add_argument("--note", help="text added to the config's notes in the row (for example the machine)")
    args = parser.parse_args()
    cfg = yaml.safe_load(open(args.config))
    if args.note:
        cfg["notes"] = "; ".join(filter(None, [cfg["notes"], args.note]))
    dice, iou, empty = evaluate(cfg["pred_dir"], Path(cfg["val_dir"]) / "masks", args.final)
    print(f"{cfg['name']}: Dice {dice:.4f}  IoU {iou:.4f}  empty {empty}  ({cfg['pred_dir']})")
    if args.save:
        save_run(cfg, args.config, dice, iou, empty, args.logs, args.date, args.hours)


if __name__ == "__main__":
    main()
